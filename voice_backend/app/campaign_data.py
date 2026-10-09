# ruff: noqa: E501  (the prompts below are long lines of prose on purpose)
"""
Data side of the telesales campaign: the customer base (from the Excel workbook), what the agent must ask on each
call, what is stored about a call, the post-call analysis and the filled workbook.

Documents (see store.py):
  campaign/base.xlsx                         the workbook as uploaded
  campaign/clients.json                      customers parsed from the sheet "Base Priorizada"
  campaign/batch.json                        ids of the customers in the current batch
  campaign/calls/<cid>/attempt.json          last call placed to that customer
  campaign/calls/<cid>/<callSid>/status.json final status reported by Twilio
  campaign/calls/<cid>/<callSid>/stream.json transcript, timing and recording id (written by the audio stream)
  campaign/results/<cid>.json                consolidated result of the last call (what goes into the Excel)
  campaign/rehearsals/<cid>.json             same, for rehearsal calls made to a test number (never exported)
"""

import io
import json
import logging
import re
import time
from datetime import date, datetime, timedelta, timezone
from typing import Any, Optional

from openai import AsyncOpenAI

from . import store
from .agent.llm_client import reasoning_options
from .config import settings

log = logging.getLogger("hypercheap.campaign")

SHEET = "Base Priorizada"
LOG_SHEET = "Llamadas IA"
OTHER = "otros"  # bucket for calls that do not belong to a campaign customer (inbound, manual tests)

ESTADOS = ["Pendiente", "Contactado", "Seguimiento", "Cotización", "Venta", "No interesado", "No localizable"]
RESULTADOS = [
    "No contesta",
    "Interesado",
    "Solicita cotización",
    "Precio alto",
    "Compra competencia",
    "Sin necesidad actual",
    "Datos incorrectos",
    "Atendido por otro vendedor",
    "Negocio cerrado",
]
CATEGORIAS = ["Baterías", "Lubricantes", "Filtros", "Aditivos", "Grasas", "Antincongelante", "Quimicos", "Otros"]
CLASIFICACIONES = [
    "Mayorista",
    "Refaccionaria",
    "Tienda Especializada",
    "Gasolinera",
    "Taller-Lubricentro",
    "Agencias Automotrices",
    "Flotillas",
    "Cadena Comercial-Lubricentro",
    "Gobierno",
    "Otros",
]

# Workbook columns (header text -> key). The last eleven are the ones a call can fill.
COLUMNS = {
    "Cliente": "id",
    "Nombre-Pagador": "nombre",
    "Oficina de ventas": "oficina",
    "Población": "poblacion",
    "Calle": "calle",
    "Compra 2026": "compra_2026",
    "Segmento piloto": "segmento",
    "Objetivo comercial": "objetivo",
    "Teléfono": "telefono",
    "Telefono 2": "telefono_2",
    "Nombre del contacto": "contacto",
    "Clasificacion Cliente": "clasificacion",
    "Causas": "causas",
    "Estado gestión": "estado_gestion",
    "Fecha último contacto": "fecha_contacto",
    "Resultado": "resultado",
    "Próximo seguimiento": "proximo_seguimiento",
    "Venta lograda ₡": "venta_lograda",
    "Categoría vendida": "categoria_vendida",
    "Observaciones": "observaciones",
}

END_MARKER = "<<FIN>>"
END_RE = re.compile(r"<<?\s*FIN\s*>>?|\[\s*FIN\s*\]", re.IGNORECASE)


def now_local() -> datetime:
    """Costa Rica has no daylight saving time, so a fixed offset is exact."""
    return datetime.now(timezone(timedelta(hours=settings.call_utc_offset)))


def within_call_hours(moment: Optional[datetime] = None) -> bool:
    moment = moment or now_local()
    try:
        start, end = (int(x) for x in settings.call_hours.split("-", 1))
    except ValueError:
        return True
    return moment.weekday() < 6 and start <= moment.hour < end  # Monday to Saturday


# --- Customers ----------------------------------------------------------------------------------------------------


def fix_text(value: Any) -> str:
    """The export has some names with broken accents (UTF-8 read as Latin-1); repair them when possible."""
    text = str(value or "").strip()
    if "Ã" in text or "Â" in text:
        try:
            return text.encode("cp1252").decode("utf-8")
        except UnicodeError:
            return text
    return text


def clean_phone(value: Any) -> str:
    """Returns the 8 digits of a Costa Rican number, or '' when the cell has no usable phone."""
    digits = re.sub(r"\D", "", str(value if value is not None else ""))
    if len(digits) == 11 and digits.startswith("506"):
        digits = digits[3:]
    return digits if len(digits) == 8 else ""


def to_e164(phone8: str) -> str:
    return f"+506{phone8}"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def parse_workbook(data: bytes) -> list[dict]:
    """Reads the customers from the workbook. Raises ValueError with a readable message when it does not fit."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError("El archivo no es un Excel válido (.xlsx).") from e
    if SHEET not in wb.sheetnames:
        raise ValueError(f'El Excel no tiene la hoja "{SHEET}".')
    rows = wb[SHEET].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    missing = [name for name in COLUMNS if name not in header]
    if missing:
        raise ValueError("Faltan columnas en la hoja: " + ", ".join(missing))
    index = {key: header.index(name) for name, key in COLUMNS.items()}

    clients: list[dict] = []
    for row_number, row in enumerate(rows, start=2):
        if row is None or index["id"] >= len(row) or row[index["id"]] in (None, ""):
            continue
        get = lambda key: row[index[key]] if index[key] < len(row) else None  # noqa: E731
        compra = get("compra_2026")
        clients.append(
            {
                "id": str(get("id")).strip(),
                "fila": row_number,
                "nombre": fix_text(get("nombre")),
                "oficina": _cell_text(get("oficina")),
                "poblacion": fix_text(get("poblacion")),
                "calle": fix_text(get("calle")),
                "compra_2026": float(compra) if isinstance(compra, (int, float)) else 0.0,
                "segmento": _cell_text(get("segmento")),
                "objetivo": _cell_text(get("objetivo")),
                "telefono": clean_phone(get("telefono")),
                "telefono_2": clean_phone(get("telefono_2")),
                "contacto": fix_text(get("contacto")),
                "clasificacion": _cell_text(get("clasificacion")),
                "causas": _cell_text(get("causas")),
                "estado_gestion": _cell_text(get("estado_gestion")) or "Pendiente",
                "fecha_contacto": _cell_text(get("fecha_contacto")),
                "resultado": _cell_text(get("resultado")),
            }
        )
    wb.close()
    if not clients:
        raise ValueError("La hoja no tiene clientes.")
    return clients


async def load_clients() -> list[dict]:
    return await store.get_json("campaign/clients.json", default=[]) or []


async def find_client(cid: str) -> Optional[dict]:
    if not cid or cid == OTHER or not store.ready():
        return None
    for client in await load_clients():
        if client["id"] == cid:
            return client
    return None


def pick_batch(clients: list[dict], segment: str, size: int, done: set[str]) -> list[str]:
    """Next customers to call: with a phone, nobody has worked them yet, and without repeating a phone number."""
    chosen: list[str] = []
    phones: set[str] = set()
    for client in clients:
        if len(chosen) >= size:
            break
        if segment and not client["segmento"].upper().startswith(segment.upper()):
            continue
        if not client["telefono"] or client["telefono"] in phones:
            continue
        worked = client["estado_gestion"] != "Pendiente" or client.get("fecha_contacto") or client.get("resultado")
        if worked or client["id"] in done:
            continue
        phones.add(client["telefono"])
        chosen.append(client["id"])
    return chosen


# --- What the agent says ------------------------------------------------------------------------------------------

_SEGMENT_GOAL = {
    "P1": (
        "Este cliente no ha comprado en dos mil veintiséis. Cuéntele que lo llama porque hace rato no sabemos de "
        "él y pregunte con interés genuino qué pasó: si fue el precio, si le compra a otro proveedor, si nadie lo "
        "visita, si ya no necesita el producto o si el negocio cambió. Escuche y no discuta."
    ),
    "P2": (
        "Este cliente sigue comprando, pero poco. Agradezca la preferencia y pregunte qué le haría falta para "
        "comprarnos más seguido: surtido, precio, visitas o entrega."
    ),
    "P3": (
        "Este cliente compra con regularidad. Agradezca la preferencia y explore qué otras líneas de Gonher podría "
        "llevar además de las que ya compra."
    ),
    "P4": (
        "Este cliente compra con regularidad. Es una llamada de atención: agradezca la preferencia, pregunte cómo le "
        "ha ido con el servicio y si necesita algo."
    ),
}


def call_prompt(client: dict) -> str:
    """Instructions appended to the base system prompt for a campaign call to this customer."""
    segment = (client.get("segmento") or "")[:2].upper()
    known = []
    if client.get("contacto"):
        known.append(f"contacto registrado: {client['contacto']}")
    if client.get("clasificacion"):
        known.append(f"tipo de negocio registrado: {client['clasificacion']}")
    if client.get("causas"):
        known.append(f"causa anotada antes: {client['causas']}")
    known_text = "; ".join(known) if known else "no hay contacto ni tipo de negocio registrados"
    kinds = ", ".join(CLASIFICACIONES[:-1]).lower()
    return f"""

LLAMADA DE TELEVENTAS, SOLO PARA ESTA LLAMADA
Usted llamó por teléfono a un cliente de la base de Gonher. Datos: cliente {client['nombre']}, de {client['poblacion'] or 'Costa Rica'}; {known_text}.
{_SEGMENT_GOAL.get(segment, _SEGMENT_GOAL['P4'])}
Además de eso, consiga estos datos con naturalidad, una sola pregunta por turno y sin que suene a encuesta:
Uno. Con quién tiene el gusto y si es la persona que ve las compras. Si no lo es, pregunte por el nombre del encargado.
Dos. A qué se dedica el negocio, para clasificarlo: {kinds}.
Tres. Qué productos usa o necesita ahora: baterías, lubricantes, filtros, aditivos, grasas, anticongelante o químicos. Si hay interés, ofrezca que un asesor le envíe la cotización.
Cuatro. Otro teléfono o WhatsApp donde contactarlo, si lo quiere dar. Repítalo para confirmarlo.
Cinco. Cuándo le conviene que lo vuelvan a contactar.
Reglas de esta llamada:
En el saludo usted ya dijo que es el asistente virtual de Gonher y que la llamada se graba. Si la persona no quiere ser grabada o prefiere hablar con una persona, ofrezca que un asesor la llame, agradezca y despídase.
Sea breve: la llamada ideal dura dos o tres minutos. Si la persona está ocupada, pregunte cuándo llamarla y despídase.
Si dice que no le interesa o pide que no la llamen más, agradezca, despídase y no insista.
Si es un número equivocado o no conocen al cliente, discúlpese y despídase.
Si contesta un buzón de voz o una grabación, no deje mensaje.
Cuando ya se haya despedido, o si contestó un buzón, termine su respuesta escribiendo exactamente {END_MARKER} para colgar. No lo escriba antes de despedirse."""


def opening_cue(client: Optional[dict], outbound: bool) -> str:
    """Stage direction that makes the agent speak first. Only the LLM sees it."""
    recorded = " y avise que la llamada se graba para calidad del servicio" if settings.call_recording else ""
    if client:
        return (
            f"(Indicación interna: usted acaba de llamar por teléfono a {client['nombre']} y ya contestaron. Salude "
            f"con calidez, diga que le habla Enrique, el asistente virtual de Gonher{recorded}, y pregunte si se "
            "comunica con ese cliente. No pregunte nada más en este turno.)"
        )
    if outbound:
        return (
            "(Indicación interna: usted acaba de llamar por teléfono a esta persona y ya contestó. Salude con "
            f"calidez, preséntese como Enrique, el asistente virtual de ventas de Gonher{recorded}, y pregunte si "
            "tiene un minuto para conversar.)"
        )
    return (
        "(Indicación interna: un cliente acaba de llamar por teléfono a Gonher. Conteste con un saludo breve y "
        f"alegre, preséntese como Enrique, el asistente virtual de ventas de Gonher{recorded}, y pregunte en qué "
        "le puede ayudar.)"
    )


# --- What is stored about a call ----------------------------------------------------------------------------------


def _safe(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]", "", str(value or ""))[:64] or OTHER


def call_dir(cid: str, call_sid: str) -> str:
    return f"campaign/calls/{_safe(cid)}/{_safe(call_sid)}"


async def save_attempt(cid: str, call_sid: str, phone: str, rehearsal: bool = False) -> None:
    """`rehearsal`: the call went to a test number, so its result must not reach the workbook."""
    doc = {
        "callSid": call_sid,
        "telefono": phone,
        "inicio": now_local().isoformat(timespec="seconds"),
        "prueba": rehearsal,
    }
    await store.put_json(f"campaign/calls/{_safe(cid)}/attempt.json", doc)


async def save_status(cid: str, call_sid: str, status: str, duration: str) -> None:
    doc = {"estado": status, "duracion_s": int(duration) if str(duration).isdigit() else 0}
    await store.put_json(f"{call_dir(cid, call_sid)}/status.json", doc)


async def save_stream(cid: str, call_sid: str, doc: dict) -> None:
    await store.put_json(f"{call_dir(cid, call_sid)}/stream.json", doc)


class Transcript:
    """Collects what was said in a call from the agent callbacks."""

    def __init__(self) -> None:
        self.lines: list[dict] = []  # {"quien": "agente" | "cliente", "texto": ..., "t": seconds}
        self.started = time.time()
        self._reply = ""

    def _add(self, who: str, text: str) -> None:
        self.lines.append({"quien": who, "texto": text, "t": round(time.time() - self.started, 1)})

    def heard(self, text: str) -> None:
        """The customer finished a sentence (this also closes a reply they interrupted)."""
        self.flush(interrupted=True)
        self._add("cliente", text)

    def token(self, tok: str) -> None:
        self._reply += tok

    def flush(self, interrupted: bool = False) -> str:
        """Closes the reply the agent was saying and returns its text ('' when there was none)."""
        text = END_RE.sub("", self._reply).strip()
        self._reply = ""
        if text:
            self._add("agente", text + (" (interrumpido)" if interrupted else ""))
        return text

    @property
    def seconds(self) -> int:
        return int(time.time() - self.started)


def transcript_text(lines: list[dict]) -> str:
    who = {"agente": "Agente", "cliente": "Cliente"}
    return "\n".join(f"{who.get(line.get('quien'), '¿?')}: {line.get('texto', '')}" for line in lines)


_NO_ANSWER = {"busy": "ocupado", "no-answer": "no contestó", "failed": "la llamada falló", "canceled": "cancelada"}

ANALYSIS_PROMPT = f"""Eres un analista de televentas. Lee la transcripción de una llamada entre el agente de Gonher y un cliente y devuelve SOLO un objeto JSON, sin texto adicional, con estas claves:
"hablo_con_persona": true si contestó una persona; false si fue buzón de voz, grabación o nadie habló.
"nombre_contacto": nombre de la persona con quien se habló o del encargado de compras, o "".
"telefono_2": otro teléfono o WhatsApp que dio el cliente, solo dígitos, o "".
"clasificacion": una de {json.dumps(CLASIFICACIONES, ensure_ascii=False)} según a qué se dedica el negocio, o "" si no se supo.
"causas": por qué dejó de comprar o compra poco, en una frase corta, o "".
"estado_gestion": una de {json.dumps(ESTADOS, ensure_ascii=False)}. Usa "Cotización" si pidió cotización, "Venta" solo si confirmó un pedido, "Seguimiento" si hay que volver a llamar, "No interesado" si lo dijo o pidió no ser llamado, "No localizable" si es número equivocado, "Contactado" en los demás casos en que se habló con una persona.
"resultado": una de {json.dumps(RESULTADOS, ensure_ascii=False)}, la que mejor describa la llamada.
"proximo_seguimiento": fecha acordada para volver a contactar en formato AAAA-MM-DD, o "". Hoy es {{hoy}}.
"categoria_vendida": si confirmó un pedido, una de {json.dumps(CATEGORIAS, ensure_ascii=False)}; si no hubo pedido, "".
"venta_lograda": monto en colones del pedido confirmado como número, o null si no hubo pedido o no se dijo el monto.
"resumen": resumen de la llamada en dos o tres frases: qué dijo el cliente, qué productos le interesan, qué quedó pendiente y si pidió no ser llamado o no ser grabado.
No inventes datos: si algo no se dijo en la llamada, déjalo vacío."""


def _pick(value: Any, options: list[str]) -> str:
    text = str(value or "").strip().lower()
    for option in options:
        if option.strip().lower() == text:
            return option
    return ""


def _parse_json(text: str) -> dict:
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        return {}
    try:
        value = json.loads(text[start : end + 1])
    except ValueError:
        return {}
    return value if isinstance(value, dict) else {}


async def analyze(lines: list[dict]) -> dict:
    """Asks the LLM to turn the transcript into the fields of the workbook. Never raises."""
    client = AsyncOpenAI(api_key=settings.baseten_api_key, base_url=settings.baseten_base_url)
    today = now_local()
    weekday = ("lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo")[today.weekday()]
    prompt = ANALYSIS_PROMPT.replace("{hoy}", f"{weekday} {today:%Y-%m-%d}")
    request: dict = dict(
        model=settings.baseten_model,
        messages=[
            {"role": "system", "content": prompt},
            {"role": "user", "content": "Transcripción:\n" + transcript_text(lines)},
        ],
        temperature=0,
    )
    raw: dict = {}
    try:
        # First a direct answer; if that fails, once more with the model's default reasoning and a larger
        # budget (hidden reasoning counts against max_tokens, which is what used to leave the JSON unfinished).
        direct = {"max_tokens": 1500, **reasoning_options(settings.baseten_reasoning_effort)}
        for options in (direct, {"max_tokens": 8000}):
            try:
                resp = await client.chat.completions.create(**request, **options)
                raw = _parse_json(resp.choices[0].message.content or "")
                if raw:
                    break
                log.warning("[campaign] analysis gave no JSON (finish=%s)", resp.choices[0].finish_reason)
            except Exception as e:
                log.error("[campaign] analysis failed: %s", e)
    finally:
        await client.close()

    if not raw:
        return {"analizado": False, "resumen": "No se pudo analizar la llamada; revise la transcripción."}
    follow_up = str(raw.get("proximo_seguimiento") or "").strip()
    sale = raw.get("venta_lograda")
    return {
        "analizado": True,
        "hablo_con_persona": bool(raw.get("hablo_con_persona", True)),
        "nombre_contacto": str(raw.get("nombre_contacto") or "").strip()[:80],
        "telefono_2": clean_phone(raw.get("telefono_2")),
        "clasificacion": _pick(raw.get("clasificacion"), CLASIFICACIONES),
        "causas": str(raw.get("causas") or "").strip()[:200],
        "estado_gestion": _pick(raw.get("estado_gestion"), ESTADOS),
        "resultado": _pick(raw.get("resultado"), RESULTADOS),
        "proximo_seguimiento": follow_up if re.fullmatch(r"\d{4}-\d{2}-\d{2}", follow_up) else "",
        "categoria_vendida": _pick(raw.get("categoria_vendida"), CATEGORIAS),
        "venta_lograda": float(sale) if isinstance(sale, (int, float)) and not isinstance(sale, bool) else None,
        "resumen": str(raw.get("resumen") or "").strip()[:900],
    }


async def call_snapshot(cid: str, force: bool = False) -> dict:
    """
    Current state of the last call to a customer. Consolidates it (and analyses it) once the call is over.
    `force` closes a call whose audio stream never wrote its final transcript.
    """
    base = f"campaign/calls/{_safe(cid)}"
    attempt = await store.get_json(f"{base}/attempt.json")
    if not attempt:
        return await store.get_json(f"campaign/results/{_safe(cid)}.json") or {"fase": "sin_llamar"}
    call_sid = attempt["callSid"]
    rehearsal = bool(attempt.get("prueba"))
    result_path = f"campaign/{'rehearsals' if rehearsal else 'results'}/{_safe(cid)}.json"
    result = await store.get_json(result_path)
    if result and result.get("callSid") == call_sid:
        return result

    folder = call_dir(cid, call_sid)
    status = await store.get_json(f"{folder}/status.json")
    stream = await store.get_json(f"{folder}/stream.json")
    started = datetime.fromisoformat(attempt["inicio"])
    waited = (now_local() - started).total_seconds()
    # Twilio's status callback is the normal end signal; the timeout covers a callback that never arrives.
    ended = bool(status) or (stream and stream.get("terminada") and waited > 90) or waited > 900
    # A browser rehearsal has no Twilio callback: if its closing save never arrived, close it with what was kept.
    ended = ended or (rehearsal and force and bool(stream))
    if not ended:
        return {"fase": "en_llamada" if stream else "marcando", "callSid": call_sid, "inicio": attempt["inicio"]}
    if status and stream and not stream.get("terminada") and waited < 900 and not force:
        # The call ended but the audio stream has not written its last transcript yet; ask again shortly.
        return {"fase": "cerrando", "callSid": call_sid, "inicio": attempt["inicio"]}

    lines = (stream or {}).get("transcripcion") or []
    spoke = any(line.get("quien") == "cliente" for line in lines)
    state = (status or {}).get("estado") or ("completed" if stream else "no-answer")
    doc: dict[str, Any] = {
        "fase": "terminada",
        "prueba": rehearsal,
        "cid": cid,
        "callSid": call_sid,
        "telefono": attempt.get("telefono", ""),
        "fecha": attempt["inicio"][:10],
        "inicio": attempt["inicio"],
        "duracion_s": (status or {}).get("duracion_s") or (stream or {}).get("duracion_s") or 0,
        "estado_llamada": state,
        "grabacion": (stream or {}).get("recordingSid", ""),
        "transcripcion": lines,
    }
    if spoke:
        doc.update(await analyze(lines))
        if not doc.get("hablo_con_persona", True):
            doc.update({"estado_gestion": "Pendiente", "resultado": "No contesta"})
        elif not doc.get("estado_gestion"):
            doc["estado_gestion"] = "Contactado"
    else:
        reason = _NO_ANSWER.get(state, "contestaron pero nadie habló" if lines else "no contestó")
        doc.update(
            {
                "analizado": True,
                "hablo_con_persona": False,
                "estado_gestion": "Pendiente",
                "resultado": "No contesta",
                "resumen": f"Intento de llamada sin contacto: {reason}.",
            }
        )
    await store.put_json(result_path, doc)
    return doc


# --- Filled workbook ----------------------------------------------------------------------------------------------


def _as_date(text: str) -> Any:
    try:
        return datetime.strptime(text, "%Y-%m-%d")
    except (TypeError, ValueError):
        return text or None


def fill_workbook(data: bytes, results: list[dict], panel_url: str) -> bytes:
    """Writes the call results into a copy of the uploaded workbook and adds a sheet with the call log."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET]
    header = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    col = {key: header[name] for name, key in COLUMNS.items() if name in header}
    rows = {str(ws.cell(row=r, column=col["id"]).value).strip(): r for r in range(2, ws.max_row + 1)}

    def put(row: int, key: str, value: Any, keep_existing: bool = False) -> None:
        if value in (None, "") or key not in col:
            return
        cell = ws.cell(row=row, column=col[key])
        if keep_existing and cell.value not in (None, ""):
            return
        cell.value = value

    for res in results:
        row = rows.get(str(res.get("cid")))
        if not row:
            continue
        put(row, "telefono_2", res.get("telefono_2"), keep_existing=True)
        put(row, "contacto", res.get("nombre_contacto"), keep_existing=True)
        put(row, "clasificacion", res.get("clasificacion"), keep_existing=True)
        put(row, "causas", res.get("causas"))
        put(row, "estado_gestion", res.get("estado_gestion"))
        put(row, "fecha_contacto", _as_date(res.get("fecha", "")))
        put(row, "resultado", res.get("resultado"))
        put(row, "proximo_seguimiento", _as_date(res.get("proximo_seguimiento", "")))
        put(row, "venta_lograda", res.get("venta_lograda"))
        put(row, "categoria_vendida", res.get("categoria_vendida"))
        put(row, "observaciones", res.get("resumen"))
        for key in ("fecha_contacto", "proximo_seguimiento"):
            if key in col:
                ws.cell(row=row, column=col[key]).number_format = "DD/MM/YYYY"

    # openpyxl drops the three drop-down lists that point at the sheet "Catálogos"; put them back.
    last = max(ws.max_row, 2)
    if "Catálogos" in wb.sheetnames:
        for key, source in (
            ("categoria_vendida", "$C$2:$C$9"),
            ("clasificacion", "$D$2:$D$11"),
            ("resultado", "$B$2:$B$10"),
        ):
            if key in col:
                letter = ws.cell(row=1, column=col[key]).column_letter
                rule = DataValidation(type="list", formula1=f"'Catálogos'!{source}", allow_blank=True)
                rule.add(f"{letter}2:{letter}{last}")
                ws.add_data_validation(rule)

    if LOG_SHEET in wb.sheetnames:
        del wb[LOG_SHEET]
    log_ws = wb.create_sheet(LOG_SHEET)
    titles = ["Fecha", "Cliente", "Nombre", "Teléfono", "Duración (s)", "Estado gestión", "Resultado", "Resumen"]
    titles += ["Transcripción", "Grabación (id)", "Escuchar en el panel"]
    log_ws.append(titles)
    for cell in log_ws[1]:
        cell.font = Font(bold=True)
    for res in sorted(results, key=lambda r: r.get("inicio", "")):
        row = rows.get(str(res.get("cid")))
        name = ws.cell(row=row, column=col["nombre"]).value if row else ""
        recording = res.get("grabacion", "")
        log_ws.append(
            [
                res.get("inicio", "").replace("T", " ")[:16],
                res.get("cid", ""),
                name,
                res.get("telefono", ""),
                res.get("duracion_s", 0),
                res.get("estado_gestion", ""),
                res.get("resultado", ""),
                res.get("resumen", ""),
                transcript_text(res.get("transcripcion") or [])[:32000],
                recording,
                f"{panel_url}?cliente={res.get('cid', '')}" if recording else "",
            ]
        )
    for letter, width in zip("ABCDEFGHIJK", (17, 11, 34, 12, 12, 16, 22, 60, 90, 38, 48)):
        log_ws.column_dimensions[letter].width = width
    for row_cells in log_ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in "HI")
    log_ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
